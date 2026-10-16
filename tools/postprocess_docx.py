"""Quarto post-render step for report.docx.

Quarto wraps every captioned figure and table into a layout table, and pandoc gives
data tables no borders. This script centers all tables and adds grid borders
(DSTU style) only to data tables, i.e. tables that hold text, not images.
"""

import os
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Emu, Pt

TABLE_FONT_SIZE = Pt(12)


def all_tables(container):
    for table in container.tables:
        yield table
        for cell in table._cells:
            yield from all_tables(cell)


def is_data_table(table):
    has_image = table._tbl.xpath(".//w:drawing")
    return not has_image and len(table.columns) > 1


def set_grid_borders(table):
    tbl_pr = table._tbl.tblPr
    old = tbl_pr.find(qn("w:tblBorders"))
    if old is not None:
        tbl_pr.remove(old)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        line = OxmlElement(f"w:{edge}")
        line.set(qn("w:val"), "single")
        line.set(qn("w:sz"), "4")
        line.set(qn("w:color"), "000000")
        borders.append(line)
    tbl_pr.append(borders)


def stretch_to_width(table, width):
    """Make a one-column layout table as wide as the text area."""
    for grid_col in table._tbl.tblGrid.findall(qn("w:gridCol")):
        grid_col.set(qn("w:w"), str(int(width.twips)))
    for cell in table._cells:
        cell.width = width


def keep_rows_on_one_page(table):
    """A figure and its caption must not be split between pages."""
    for row in table.rows:
        tr_pr = row._tr.get_or_add_trPr()
        if tr_pr.find(qn("w:cantSplit")) is None:
            tr_pr.append(OxmlElement("w:cantSplit"))


CAPTION_STYLES = {"ImageCaption", "TableCaption"}
TABLE_CAPTION_PREFIX = "Таблиця"  # crossref tbl-title in report.qmd


def fix_captions(doc):
    """Quarto writes captions with two <w:pPr>; Word applies the last one (left).

    Keep only the properties block that holds the caption style and align it:
    figure captions centered, table captions left (DSTU).
    """
    for p in doc.element.body.iter(qn("w:p")):
        props = p.findall(qn("w:pPr"))
        if not props:
            continue
        styled = [pr for pr in props if pr.find(qn("w:pStyle")) is not None]
        keep = styled[-1] if styled else props[-1]
        for pr in props:
            if pr is not keep:
                p.remove(pr)

        style = keep.find(qn("w:pStyle"))
        if style is None or style.get(qn("w:val")) not in CAPTION_STYLES:
            continue
        text = "".join(t.text or "" for t in p.iter(qn("w:t")))
        alignment = "left" if text.startswith(TABLE_CAPTION_PREFIX) else "center"
        jc = keep.find(qn("w:jc"))
        if jc is None:
            jc = OxmlElement("w:jc")
            keep.append(jc)
        jc.set(qn("w:val"), alignment)


def process(path):
    doc = Document(path)
    fix_captions(doc)
    section = doc.sections[0]
    text_width = Emu(section.page_width - section.left_margin - section.right_margin)

    for table in all_tables(doc):
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        if len(table.columns) == 1:
            stretch_to_width(table, text_width)
            keep_rows_on_one_page(table)
        if is_data_table(table):
            set_grid_borders(table)
            for cell in table._cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.size = TABLE_FONT_SIZE
    doc.save(path)


if __name__ == "__main__":
    outputs = (
        os.environ.get("QUARTO_PROJECT_OUTPUT_FILES", "").splitlines() or sys.argv[1:]
    )
    for output in outputs:
        if output.endswith(".docx"):
            process(output)
